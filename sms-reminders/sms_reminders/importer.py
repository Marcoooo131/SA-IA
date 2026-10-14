"""CSV / Excel import: parsing, column mapping, validation, preview, commit."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .appointments import AppointmentData, business_key, create_appointment
from .audit import record
from .models import Appointment, ImportBatch, Service, utcnow
from .phone import check_mobile
from .policy import Policy
from .timeutil import LocalTimeError, local_to_utc


class ImportFileError(ValueError):
    pass


FIELDS: dict[str, str] = {
    "customer_ref": "Numéro client interne",
    "full_name": "Nom et prénom (une seule colonne)",
    "last_name": "Nom",
    "first_name": "Prénom",
    "phone": "Téléphone mobile",
    "date": "Date du rendez-vous",
    "time": "Heure du rendez-vous",
    "datetime": "Date et heure (une seule colonne)",
    "service": "Service (Movers / Takeovers)",
    "confirmed": "Rendez-vous confirmé (facultatif)",
    "external_ref": "Identifiant du rendez-vous dans le CRM (facultatif)",
}

_SYNONYMS = {
    "customer_ref": ["numero client", "numero client interne", "n client", "no client", "num client", "id client",
                     "client id", "reference client", "ref client", "code client", "customer id", "customer ref"],
    "external_ref": ["id rendez vous", "identifiant rendez vous", "identifiant rdv", "id rdv", "ref rdv",
                     "numero rdv", "numero rendez vous", "appointment id"],
    "first_name": ["prenom", "first name", "firstname"],
    "last_name": ["nom", "nom de famille", "last name", "lastname"],
    "full_name": ["nom et prenom", "nom prenom", "prenom nom", "nom complet", "nom du client", "client", "full name"],
    "phone": ["telephone", "tel", "mobile", "portable", "numero de telephone", "numero telephone", "phone", "gsm",
              "telephone mobile"],
    "date": ["date", "date rdv", "date du rendez vous", "date rendez vous", "jour"],
    "time": ["heure", "heure rdv", "heure du rendez vous", "heure rendez vous", "horaire"],
    "datetime": ["date et heure", "date heure", "datetime", "date heure rdv"],
    "service": ["service", "activite", "pole", "type de service"],
    "confirmed": ["confirme", "confirmation", "statut", "statut confirmation", "rendez vous confirme", "rdv confirme"],
}


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def guess_mapping(headers: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {k: None for k in FIELDS}
    used: set[str] = set()
    normed = {h: _norm(h) for h in headers}
    for field_key, synonyms in _SYNONYMS.items():
        for h in headers:
            if h not in used and normed[h] in synonyms:
                mapping[field_key] = h
                used.add(h)
                break
    if mapping["first_name"] and mapping["last_name"]:
        mapping["full_name"] = None
    return mapping


# ------------------------------------------------------------------ file parsing


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_file(filename: str, content: bytes, max_rows: int) -> tuple[str, list[str], list[list[str]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        origin, headers, rows = "xlsx", *_parse_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        origin, headers, rows = "csv", *_parse_csv(content)
    else:
        raise ImportFileError("Format non pris en charge : utilisez un fichier .csv ou .xlsx.")
    if not headers or not any(headers):
        raise ImportFileError("La première ligne doit contenir les noms de colonnes.")
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ImportFileError("Le fichier ne contient aucune ligne de données.")
    if len(rows) > max_rows:
        raise ImportFileError(f"Trop de lignes ({len(rows)}) : maximum {max_rows} par import.")
    headers = [h or f"Colonne {i + 1}" for i, h in enumerate(headers)]
    if len(set(headers)) != len(headers):
        raise ImportFileError("Plusieurs colonnes portent le même nom : renommez-les avant l'import.")
    width = len(headers)
    rows = [(r + [""] * width)[:width] for r in rows]
    return origin, headers, rows


def _parse_csv(content: bytes) -> tuple[list[str], list[list[str]]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise ImportFileError("Encodage du fichier non reconnu (utilisez UTF-8).")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    all_rows = [[c.strip() for c in row] for row in reader]
    if not all_rows:
        return [], []
    return all_rows[0], all_rows[1:]


def _parse_xlsx(content: bytes) -> tuple[list[str], list[list[str]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parsing failure is a user error here
        raise ImportFileError("Fichier Excel illisible.") from exc
    ws = wb.worksheets[0]
    rows = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        return [], []
    return rows[0], rows[1:]


# ------------------------------------------------------------------ value parsing


_DATE_FORMATS = ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d")


def parse_date(raw: str) -> date:
    raw = raw.strip()
    if not raw:
        raise ValueError("date manquante")
    candidate = raw.split(" ")[0].split("T")[0]
    for f in _DATE_FORMATS:
        try:
            return datetime.strptime(candidate, f).date()
        except ValueError:
            continue
    if re.fullmatch(r"\d{5}(\.0+)?", raw):  # Excel serial number exported as text
        return date(1899, 12, 30) + timedelta(days=int(float(raw)))
    raise ValueError(f"date illisible « {raw} » (attendu JJ/MM/AAAA)")


def parse_time(raw: str) -> time:
    raw = raw.strip().lower().replace(" ", "")
    if not raw:
        raise ValueError("heure manquante")
    m = re.fullmatch(r"(\d{1,2})(?:[:h](\d{2})?)?(?::(\d{2}))?", raw)
    if m:
        h, mi = int(m.group(1)), int(m.group(2) or 0)
        if h < 24 and mi < 60:
            return time(h, mi)
    if re.fullmatch(r"0?\.\d+", raw):  # Excel fraction of a day
        minutes = round(float(raw) * 24 * 60)
        return time(minutes // 60 % 24, minutes % 60)
    raise ValueError(f"heure illisible « {raw} » (attendu HH:MM)")


def parse_datetime(raw: str) -> tuple[date, time]:
    raw = raw.strip()
    parts = re.split(r"[ T]+", raw, maxsplit=1)
    if len(parts) != 2:
        raise ValueError(f"date et heure illisibles « {raw} »")
    return parse_date(parts[0]), parse_time(parts[1])


def parse_service(raw: str) -> str | None:
    n = _norm(raw).replace(" ", "")
    if not n:
        return None
    if "mover" in n:
        return Service.MOVERS
    if "takeover" in n:
        return Service.TAKEOVERS
    raise ValueError(f"service inconnu « {raw} » (attendu Movers ou Takeovers)")


_TRUE = {"oui", "o", "yes", "y", "1", "true", "vrai", "x", "confirme", "confirmee"}
_FALSE = {"", "non", "n", "no", "0", "false", "faux", "a confirmer", "non confirme", "en attente"}


def parse_confirmed(raw: str) -> tuple[bool, str | None]:
    n = _norm(raw)
    if n in _TRUE:
        return True, None
    if n in _FALSE:
        return False, None
    return False, f"statut de confirmation « {raw} » non reconnu : considéré « à confirmer »"


def split_full_name(raw: str) -> tuple[str, str]:
    """'DUPONT Marie' / 'Marie DUPONT' / 'Marie Dupont' -> (first, last)."""
    tokens = raw.split()
    upper = [t for t in tokens if len(t) > 1 and t.isupper()]
    if upper and len(upper) < len(tokens):
        return " ".join(t for t in tokens if t not in upper), " ".join(upper)
    if len(tokens) >= 2:
        return tokens[0], " ".join(tokens[1:])
    return "", raw.strip()


def _excel_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if raw.strip().isdigit() and len(digits) == 9 and digits[0] in "67":
        return "0" + digits  # leading zero lost by Excel
    return raw


# ------------------------------------------------------------------ validation / preview


@dataclass
class RowResult:
    line: int  # line number in the file (header = 1)
    data: AppointmentData | None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    status: str = "new"  # new | duplicate_existing | duplicate_file | error
    key: str | None = None


STATUS_LABELS = {
    "new": "Nouveau",
    "duplicate_existing": "Déjà enregistré (ignoré)",
    "duplicate_file": "Doublon dans le fichier (ignoré)",
    "error": "Erreur (ignoré)",
}


def mapping_errors(mapping: dict[str, str | None]) -> list[str]:
    errors = []
    if not mapping.get("customer_ref"):
        errors.append("Associez la colonne « Numéro client interne ».")
    if not mapping.get("phone"):
        errors.append("Associez la colonne « Téléphone mobile ».")
    if not (mapping.get("full_name") or (mapping.get("first_name") and mapping.get("last_name")) or mapping.get("last_name")):
        errors.append("Associez « Nom et prénom » ou les colonnes « Nom » et « Prénom ».")
    if not (mapping.get("datetime") or (mapping.get("date") and mapping.get("time"))):
        errors.append("Associez « Date » et « Heure » (ou une colonne « Date et heure »).")
    used = [v for v in mapping.values() if v]
    if len(used) != len(set(used)):
        errors.append("Une même colonne est associée à plusieurs champs.")
    if mapping.get("phone") and mapping.get("phone") == mapping.get("customer_ref"):
        errors.append("Le numéro client et le téléphone doivent être deux colonnes différentes.")
    return errors


def validate_rows(session: Session, headers: list[str], rows: list[list[str]], mapping: dict[str, str | None], now: datetime) -> list[RowResult]:
    idx = {h: i for i, h in enumerate(headers)}

    def get(row: list[str], key: str) -> str:
        col = mapping.get(key)
        return row[idx[col]].strip() if col and col in idx else ""

    results: list[RowResult] = []
    seen: dict[str, int] = {}
    for n, row in enumerate(rows, start=2):
        errors: list[str] = []
        warnings: list[str] = []
        customer_ref = get(row, "customer_ref")
        if not customer_ref:
            errors.append("numéro client manquant")
        if mapping.get("first_name") or mapping.get("last_name"):
            first, last = get(row, "first_name"), get(row, "last_name")
        else:
            first, last = split_full_name(get(row, "full_name"))
            if first and last:
                warnings.append(f"nom découpé en prénom « {first} » et nom « {last} » : à vérifier")
        if not (first or last):
            errors.append("nom du client manquant")
        elif not first:
            warnings.append("prénom manquant : un template utilisant {{prenom}} sera bloqué")
        phone_raw = _excel_phone(get(row, "phone"))
        phone = check_mobile(phone_raw)
        if not phone.valid:
            errors.append(phone.error or "téléphone invalide")
        ref_digits = re.sub(r"\D", "", customer_ref)
        if phone.e164 and len(ref_digits) >= 9 and ref_digits.lstrip("0") in (phone.e164[3:], phone.e164.lstrip("+")):
            warnings.append("le numéro client ressemble au numéro de téléphone : vérifier la correspondance des colonnes")
        d = t = None
        try:
            if mapping.get("datetime"):
                d, t = parse_datetime(get(row, "datetime"))
            else:
                d, t = parse_date(get(row, "date")), parse_time(get(row, "time"))
        except ValueError as exc:
            errors.append(str(exc))
        service = None
        try:
            service = parse_service(get(row, "service")) if mapping.get("service") else None
        except ValueError as exc:
            errors.append(str(exc))
        if service is None and not any("service" in e for e in errors):
            warnings.append("service non renseigné : aucune relance ne sera programmée tant qu'il n'est pas défini")
        confirmed, conf_warning = parse_confirmed(get(row, "confirmed")) if mapping.get("confirmed") else (False, None)
        if conf_warning:
            warnings.append(conf_warning)
        external_ref = get(row, "external_ref") or None
        at_utc = None
        if d and t:
            try:
                at_utc, dst_warning = local_to_utc(d, t)
                if dst_warning:
                    warnings.append(dst_warning)
            except LocalTimeError as exc:
                errors.append(str(exc))
        if at_utc and at_utc <= now:
            errors.append("rendez-vous déjà passé")
        elif at_utc and at_utc - now < timedelta(hours=48) and not confirmed:
            warnings.append("rendez-vous dans moins de 48 h : les relances dont l'échéance est passée ne seront pas envoyées")

        if errors:
            results.append(RowResult(n, None, errors, warnings, "error"))
            continue
        data = AppointmentData(customer_ref, first, last, phone_raw, d, t, service, confirmed, external_ref)
        key = business_key(external_ref, customer_ref, service, d, t)
        res = RowResult(n, data, errors, warnings, key=key)
        if key in seen:
            res.status = "duplicate_file"
            res.warnings.append(f"même rendez-vous que la ligne {seen[key]}")
        else:
            seen[key] = n
            if session.scalar(select(Appointment.id).where(Appointment.business_key == key)):
                res.status = "duplicate_existing"
            else:
                others = session.scalar(select(func.count(Appointment.id)).where(Appointment.customer_ref == customer_ref.strip())) or 0
                if others:
                    res.warnings.append(f"ce client a déjà {others} autre(s) rendez-vous : vérifier s'il s'agit d'un report")
        results.append(res)
    return results


def summarize(results: list[RowResult]) -> dict[str, int]:
    out = {k: 0 for k in STATUS_LABELS}
    for r in results:
        out[r.status] += 1
    out["warnings"] = sum(1 for r in results if r.warnings and r.status == "new")
    out["total"] = len(results)
    return out


# ------------------------------------------------------------------ batch lifecycle


def create_batch(session: Session, *, filename: str, content: bytes, actor: str, max_rows: int) -> ImportBatch:
    origin, headers, rows = parse_file(filename, content, max_rows)
    batch = ImportBatch(created_by=actor, filename=filename[:255], origin=origin, headers=headers, rows=rows,
                        mapping=guess_mapping(headers))
    session.add(batch)
    session.flush()
    record(session, actor, "import_uploaded", f"Fichier importé pour vérification ({len(rows)} lignes, {origin}).",
           details={"batch_id": batch.id})
    return batch


def commit_batch(session: Session, batch: ImportBatch, *, actor: str, policy: Policy, now: datetime) -> dict[str, int]:
    if batch.status != "uploaded" or batch.rows is None:
        raise ImportFileError("Cet import a déjà été traité.")
    results = validate_rows(session, batch.headers, batch.rows, batch.mapping or {}, now)
    created = 0
    for r in results:
        if r.status != "new" or r.data is None:
            continue
        create_appointment(session, r.data, actor=actor, policy=policy, now=now,
                           origin=f"import {batch.origin} #{batch.id}", batch_id=batch.id)
        created += 1
    summary = summarize(results)
    summary["created"] = created
    batch.summary = summary
    batch.status = "committed"
    batch.committed_at = now
    batch.rows = None  # raw personal data no longer needed
    record(session, actor, "import_committed",
           f"Import #{batch.id} validé : {created} rendez-vous créés, {summary['duplicate_existing']} déjà enregistrés, "
           f"{summary['duplicate_file']} doublons, {summary['error']} erreurs.", details={"batch_id": batch.id, **summary})
    return summary


def discard_batch(session: Session, batch: ImportBatch, *, actor: str) -> None:
    batch.status = "discarded"
    batch.rows = None
    record(session, actor, "import_discarded", f"Import #{batch.id} abandonné.", details={"batch_id": batch.id})


def purge_stale_batches(session: Session, hours: int) -> int:
    cutoff = utcnow() - timedelta(hours=hours)
    stale = session.scalars(select(ImportBatch).where(ImportBatch.rows.is_not(None), ImportBatch.created_at < cutoff)).all()
    for b in stale:
        b.rows = None
        if b.status == "uploaded":
            b.status = "discarded"
    return len(stale)
